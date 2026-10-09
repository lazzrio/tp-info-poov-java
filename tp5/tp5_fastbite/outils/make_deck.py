# -*- coding: utf-8 -*-
"""Construit la presentation TP5 (10 slides) avec python-pptx -> ../../rendu_tp5/"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image

RENDU = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..", "rendu_tp5"))
CAP = os.path.join(RENDU, "captures")
UML = os.path.join(RENDU, "uml")
NOMS = "[Prénom NOM 1] · [Prénom NOM 2] · [Prénom NOM 3]"
OUTFILE = os.path.join(RENDU, "TP5_FastBite_NOM1_NOM2_NOM3.pptx")

NAVY = RGBColor(0x1F, 0x3A, 0x5F)
ORANGE = RGBColor(0xE8, 0x5D, 0x04)
GREY = RGBColor(0x55, 0x5B, 0x66)
LIGHT = RGBColor(0xF3, 0xF5, 0xF9)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK = RGBColor(0x1B, 0x1F, 0x27)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def rect(slide, x, y, w, h, fill, line=None):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
    s.shadow.inherit = False
    return s


def text(slide, x, y, w, h, paras, anchor=MSO_ANCHOR.TOP):
    """paras = [(texte, taille, gras, couleur, puce?)] ; retourne la zone de texte."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    first = True
    for p in paras:
        t, size, bold, color, bullet = (list(p) + [None] * 5)[:5]
        para = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        run = para.add_run()
        run.text = ("• " + t) if bullet else t
        run.font.size = Pt(size or 14)
        run.font.bold = bool(bold)
        run.font.name = "Calibri"
        run.font.color.rgb = color or DARK
        para.space_after = Pt(5)
    return tb


def header(slide, title, n):
    rect(slide, 0, 0, 13.333, 0.95, NAVY)
    rect(slide, 0, 0.95, 13.333, 0.06, ORANGE)
    text(slide, 0.45, 0.12, 10.6, 0.75, [(title, 27, True, WHITE)], MSO_ANCHOR.MIDDLE)
    text(slide, 10.6, 0.12, 2.4, 0.75, [("FastBite · TP5", 14, False, RGBColor(0xC9, 0xD6, 0xEA))],
         MSO_ANCHOR.MIDDLE).text_frame.paragraphs[0].alignment = PP_ALIGN.RIGHT
    text(slide, 0.45, 7.12, 9, 0.3, [("POO Java — Exceptions · Fichier · Sérialisation", 10, False, GREY)])
    t = text(slide, 11.9, 7.12, 1.0, 0.3, [(str(n), 11, True, GREY)])
    t.text_frame.paragraphs[0].alignment = PP_ALIGN.RIGHT


def picture(slide, name, x, y, w, h, caption=None, folder=CAP):
    """Place l'image dans la boite (x,y,w,h) en gardant le ratio ; legende dessous. Retourne le bas."""
    path = os.path.join(folder, name)
    iw, ih = Image.open(path).size
    cap_h = 0.3 if caption else 0
    scale = min(w / iw, (h - cap_h) / ih)
    pw, ph = iw * scale, ih * scale
    px = x + (w - pw) / 2
    slide.shapes.add_picture(path, Inches(px), Inches(y), Inches(pw), Inches(ph))
    rect_line = slide.shapes[-1].line
    rect_line.color.rgb = RGBColor(0xC4, 0xCA, 0xD4)
    rect_line.width = Pt(0.75)
    if caption:
        t = text(slide, x, y + ph + 0.02, w, 0.3, [(caption, 10, False, GREY)])
        t.text_frame.paragraphs[0].runs[0].font.italic = True
        t.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    return y + ph + cap_h


def notes(slide, s):
    slide.notes_slide.notes_text_frame.text = s


# ====================================================================== 1. Couverture
s = prs.slides.add_slide(BLANK)
rect(s, 0, 0, 13.333, 7.5, NAVY)
rect(s, 0, 4.55, 13.333, 0.08, ORANGE)
text(s, 0.9, 1.5, 11.5, 0.8, [("FastBite", 54, True, WHITE)])
text(s, 0.9, 2.55, 11.5, 0.6, [("(notre application de livraison de repas, type Deliveroo)", 20, False, RGBColor(0xC9, 0xD6, 0xEA))])
text(s, 0.9, 3.35, 11.8, 1.0, [("TP5 — Exceptions, fichier texte, sérialisation / désérialisation", 30, True, ORANGE)])
text(s, 0.9, 4.9, 11.5, 0.5, [(NOMS, 22, True, WHITE)])
text(s, 0.9, 5.5, 11.5, 1.0, [
    ("Programmation orientée objet en Java — ECE Paris, Ing2", 16, False, RGBColor(0xC9, 0xD6, 0xEA)),
    ("Enseignant : J.-P. Segado · 02/10/2026", 16, False, RGBColor(0xC9, 0xD6, 0xEA))])
notes(s, "Présentation (10 s). Notre thème libre est FastBite, une application de livraison de repas comme Deliveroo, "
         "que nous avons construite aux TP2, TP3 et TP4. Ce TP5 ajoute trois choses : la gestion des erreurs avec les exceptions, "
         "l'écriture et la lecture d'un fichier texte, et la sérialisation d'objets dans un fichier binaire.")

# ====================================================================== 2. Sommaire
s = prs.slides.add_slide(BLANK)
header(s, "Sommaire", 2)
items = [("1", "Q1 — Exceptions", "repérer, lever, propager, attraper + UML mis à jour (3 slides)"),
         ("2", "Q2 — Fichier texte", "écrire, lire et afficher le menu d'un restaurant (2 slides)"),
         ("3", "Q3 — Sérialisation", "enregistrer et relire des restaurants dans un .ser (2 slides)"),
         ("4", "Bilan, sources et IA", "ce qui marche, ce qui a changé, outils utilisés")]
for i, (n, t, d) in enumerate(items):
    y = 1.55 + i * 1.15
    c = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.9), Inches(y), Inches(0.75), Inches(0.75))
    c.fill.solid(); c.fill.fore_color.rgb = ORANGE; c.line.fill.background()
    c.text_frame.text = n
    c.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    c.text_frame.paragraphs[0].runs[0].font.size = Pt(26)
    c.text_frame.paragraphs[0].runs[0].font.bold = True
    c.text_frame.paragraphs[0].runs[0].font.color.rgb = WHITE
    text(s, 1.95, y - 0.02, 10, 0.5, [(t, 26, True, NAVY)])
    text(s, 1.95, y + 0.42, 10, 0.4, [(d, 16, False, GREY)])
rect(s, 0.9, 6.15, 11.5, 0.7, LIGHT)
text(s, 1.05, 6.2, 11.2, 0.6, [("Point de départ : notre TP4 complet (Evaluable, Utilisateur abstraite, Panier extends Commande implements "
                                 "Comparable). Le TP5 le modifie le moins possible.", 14, False, DARK)], MSO_ANCHOR.MIDDLE)
notes(s, "Sommaire (10 s). Trois questions puis un bilan avec les sources. Nous sommes partis du code exact de notre TP4 "
         "et nous l'avons modifié le moins possible, pour que les changements du TP5 restent lisibles.")

# ====================================================================== 3. Q1 table + justification
s = prs.slides.add_slide(BLANK)
header(s, "Q1 — Exceptions : repérer, lever, propager", 3)
rows = [("Exception (extends FastBiteException)", "Levée par", "Pourquoi"),
        ("NoteInvalideException", "ajouterNote (Livreur, Plat, Restaurant) ; propagée par Client.noterRider", "note hors 0–5, NaN"),
        ("AucuneNoteException", "getNoteMoyenne", "aucune note (le TP4 renvoyait 0.0)"),
        ("PaiementImpossibleException", "Client.payerCommande", "commande vide ou déjà payée"),
        ("CapaciteDepasseeException", "Client.passerCommande, Restaurant.gererMenu", "tableaux de 20 pleins"),
        ("LivreurIndisponibleException", "Livreur.accepterLivraison", "livreur déjà occupé"),
        ("PlatIndisponibleException", "Commande.ajouterPlat (et Panier)", "plat absent du menu"),
        ("CommandeVideException", "Commande.valider", "aucun plat"),
        ("EtatLivraisonException", "Livraison.demarrer / terminer", "étape dans le mauvais ordre"),
        ("FormatFichierException", "Restaurant.chargerMenu (Q2)", "ligne du fichier mal formée")]
tbl = s.shapes.add_table(len(rows), 3, Inches(0.4), Inches(1.25), Inches(7.6), Inches(4.6)).table
for j, w in enumerate((2.45, 3.25, 1.9)):
    tbl.columns[j].width = Inches(w)
for i, r in enumerate(rows):
    tbl.rows[i].height = Inches(0.44)
    for j, v in enumerate(r):
        c = tbl.cell(i, j)
        c.text = v
        c.margin_left = c.margin_right = Inches(0.06)
        c.margin_top = c.margin_bottom = Inches(0.02)
        p = c.text_frame.paragraphs[0]
        p.runs[0].font.size = Pt(10.5 if i else 11)
        p.runs[0].font.name = "Calibri"
        p.runs[0].font.bold = (i == 0 or j == 0)
        p.runs[0].font.color.rgb = WHITE if i == 0 else DARK
        c.fill.solid()
        c.fill.fore_color.rgb = NAVY if i == 0 else (LIGHT if i % 2 else WHITE)
        c.vertical_anchor = MSO_ANCHOR.MIDDLE
text(s, 0.4, 5.95, 7.6, 1.1, [
    ("Non vérifiée, conservée du TP4 : IllegalArgumentException (plat null) — c'est un bug du code appelant, "
     "pas une situation à « gérer ».", 12.5, False, GREY, True),
    ("Rien n'est attrapé dans les méthodes métier : elles lèvent ou propagent (throws), le Main attrape.", 12.5, False, GREY, True)])
bottom = picture(s, "q1_code_hierarchie.png", 8.2, 1.2, 4.8, 3.6,
                 "Capture 1 — la hiérarchie et l'interface Evaluable avec throws")
text(s, 8.2, bottom + 0.05, 4.8, 2.0, [
    ("Pourquoi vérifiées (checked) ?", 14, True, NAVY),
    ("Erreurs prévisibles et récupérables : le compilateur oblige à les traiter.", 12.5, False, DARK, True),
    ("Une classe mère FastBiteException : un seul catch suffit pour toutes les erreurs métier.", 12.5, False, DARK, True),
    ("Changement TP4 : IllegalArgumentException → NoteInvalideException ; moyenne 0.0 → AucuneNoteException.", 12.5, False, DARK, True)])
notes(s, "Q1 (1 min 15). Nous avons relu chaque classe du TP4 en cherchant ce qui peut mal se passer. Neuf cas ressortent, "
         "et chacun a sa propre exception qui hérite de FastBiteException. Ce sont des exceptions vérifiées : ce sont des erreurs prévisibles, "
         "le compilateur nous oblige donc à les traiter. Deux ont remplacé un comportement du TP4 : une note invalide levait une "
         "IllegalArgumentException, et une moyenne sans note renvoyait 0.0, ce qui se confondait avec une vraie note de zéro. "
         "Nous avons gardé IllegalArgumentException pour le plat null, car c'est un bug du programmeur, pas un cas à gérer. "
         "Les méthodes métier ne font aucun try/catch : elles lèvent ou propagent, et c'est le Main qui attrape.")

# ====================================================================== 4. Q1 UML
s = prs.slides.add_slide(BLANK)
header(s, "Q1 — UML mis à jour (base : notre diagramme TP4)", 4)
picture(s, "TP5_Q1_exceptions.png", 0.25, 1.15, 9.7, 5.9, "Capture 2 — notre diagramme TP4 complété : « throws » dans les signatures + hiérarchie des exceptions", folder=UML)
picture(s, "TP4_Q2_base_ancien_diagramme.png", 10.1, 1.2, 2.95, 2.3, "Avant : diagramme TP4 (base)", folder=UML)
text(s, 10.05, 3.65, 3.1, 3.4, [
    ("Ce qui change", 14, True, NAVY),
    ("« throws X » rouge : exceptions que la méthode peut lever.", 11, False, DARK, True),
    ("Colonne orange : FastBiteException + 9 sous-classes.", 11, False, DARK, True),
    ("Panier (TP4 Q3) ajouté ; il redéclare le throws de Commande.", 11, False, DARK, True),
    ("Signatures alignées sur le code réel.", 11, False, DARK, True)])
notes(s, "Q1, diagramme (45 s). Nous sommes partis de notre diagramme du TP4, en petit en haut à droite, avec les mêmes couleurs et les mêmes "
         "associations. Nous avons ajouté, en rouge, le mot throws dans les signatures des méthodes qui peuvent lever une exception, "
         "et la hiérarchie des exceptions à droite : FastBiteException hérite d'Exception, et les neuf exceptions héritent de FastBiteException. "
         "Nous avons aussi ajouté Panier, qui hérite de Commande : quand elle redéfinit ajouterPlat, elle doit redéclarer le même throws, "
         "sinon le compilateur refuse. Le diagramme sert de contrat : l'appelant sait ce qu'il doit gérer sans lire le code.")

# ====================================================================== 5. Q1 code + console
s = prs.slides.add_slide(BLANK)
header(s, "Q1 — Lever, propager, attraper : code et résultat", 5)
picture(s, "q1_code_throw.png", 0.3, 1.2, 6.4, 5.8, "Capture 3 — throw / throws dans Livreur et Commande")
picture(s, "q1_console_B.png", 6.9, 1.2, 6.15, 5.8, "Capture 4 — chaque exception est levée puis attrapée (B1 à B8) ; finally exécuté")
notes(s, "Q1, code et résultat (1 min). À gauche, throw dans le code : accepterLivraison lève LivreurIndisponibleException "
         "au lieu d'afficher un message, ajouterPlat vérifie que le plat est au menu, valider refuse une commande vide. "
         "À droite la console : B1 à B8, un cas d'erreur par exception, attrapé avec try/catch. Le cas B5 montre aussi le bloc finally, "
         "exécuté dans tous les cas. Le scénario complet du TP3 et du TP4, qui n'a aucun try/catch, passe sans erreur : "
         "les exceptions sont propagées jusqu'au Main, qui les attrape en un seul catch sur la classe mère.")

# ====================================================================== 6. Q2 code
s = prs.slides.add_slide(BLANK)
header(s, "Q2 — Écrire et lire un fichier texte", 6)
rect(s, 0.4, 1.25, 5.9, 1.55, LIGHT)
text(s, 0.5, 1.28, 5.7, 1.5, [
    ("Format du fichier (menu_chezmario.txt)", 14, True, NAVY),
    ("1 ligne = 1 plat, 4 champs séparés par « ; »", 12.5, False, DARK, True),
    ("nom;prix;description;allergène1,allergène2", 12.5, True, ORANGE, True),
    ("prix avec point décimal ; allergènes séparés par « , » (vide = aucun)", 12.5, False, DARK, True)])
picture(s, "q2_code_ecriture.png", 0.3, 2.9, 6.1, 2.7, "Capture 5 — écriture : sauvegarderMenu (try-with-resources)")
text(s, 0.4, 5.75, 6.0, 1.3, [
    ("Pourquoi ces choix ?", 13.5, True, NAVY),
    ("try-with-resources : le fichier est fermé même en cas d'erreur.", 11.5, False, DARK, True),
    ("IOException propagée (throws) : l'appelant décide quoi faire.", 11.5, False, DARK, True)])
picture(s, "q2_code_lecture.png", 6.5, 1.2, 6.55, 5.85, "Capture 6 — lecture : chargerMenu (split, FormatFichierException)")
notes(s, "Q2, code (1 min). Nous avons choisi un format simple : une ligne par plat, quatre champs séparés par un point-virgule, "
         "et les allergènes séparés par des virgules. À gauche l'écriture avec un try-with-resources qui ferme le fichier tout seul. "
         "À droite la lecture : split avec -1 pour garder le dernier champ vide, vérification du nombre de champs, et conversion du prix. "
         "Si la ligne est mauvaise on lève FormatFichierException avec le numéro de ligne, et pour le prix on garde l'erreur d'origine "
         "NumberFormatException comme cause : c'est le chaînage. Les plats sont d'abord lus dans une liste temporaire, "
         "donc si une ligne est fausse le menu n'est pas modifié du tout.")

# ====================================================================== 7. Q2 UML + résultats
s = prs.slides.add_slide(BLANK)
header(s, "Q2 — Résultat et UML mis à jour", 7)
b = picture(s, "TP5_Q2_fichier.png", 0.25, 1.15, 7.0, 4.35, "Capture 7 — diagramme TP4 + sauvegarderMenu / chargerMenu / afficherMenu, IOException et FormatFichierException", folder=UML)
picture(s, "q2_fichier_txt.png", 0.25, b + 0.08, 7.0, 1.45, "Capture 8 — le fichier produit (menu_chezmario.txt)")
picture(s, "q2_console.png", 7.4, 1.2, 5.65, 5.85, "Capture 9 — écriture, relecture, affichage, puis 3 erreurs")
notes(s, "Q2, résultat (1 min). En haut à gauche notre diagramme TP4 mis à jour : Restaurant a trois nouvelles méthodes, sauvegarderMenu qui peut lever "
         "IOException, chargerMenu qui peut lever aussi FormatFichierException, et afficherMenu. En dessous, le vrai fichier produit. "
         "À droite la console : on écrit le menu, on le relit dans un nouveau restaurant et on l'affiche, les trois plats sont identiques. "
         "Ensuite trois erreurs : fichier introuvable, prix non numérique avec la cause chaînée, et ligne avec deux champs au lieu de quatre. "
         "Après ces échecs le menu reste intact.")

# ====================================================================== 8. Q3 code
s = prs.slides.add_slide(BLANK)
header(s, "Q3 — Sérialiser des restaurants", 8)
picture(s, "q3_code_serializable.png", 0.3, 1.2, 6.2, 3.2, "Capture 10 — Serializable, serialVersionUID, champ transient")
text(s, 0.4, 4.55, 6.1, 2.5, [
    ("Pourquoi ces choix ?", 14, True, NAVY),
    ("Serializable sur Restaurant ET sur Plat : un objet n'est sérialisable que si tout ce qu'il contient l'est.", 12, False, DARK, True),
    ("serialVersionUID : fixe la version de la classe pour la relecture.", 12, False, DARK, True),
    ("transient : le compteur de session n'est pas enregistré (0 après relecture).", 12, False, DARK, True),
    ("Client, Livreur, Commande… non sérialisables : ce sont des sessions, pas des données.", 12, False, DARK, True)])
b = picture(s, "q3_code_ser.png", 6.7, 1.2, 6.35, 3.4, "Capture 11 — writeObject / readObject (cast + ClassNotFoundException)")
picture(s, "q3_fichier_ser.png", 6.9, b + 0.08, 5.6, 2.3, "Capture 12 — le fichier .ser (binaire, signature ac ed 00 05)")
notes(s, "Q3, code (1 min). Dans notre thème, ce que le professeur appelle les hébergements, ce sont les restaurants. "
         "Nous sérialisons une liste de restaurants. Pour cela Restaurant et Plat implémentent Serializable, parce qu'un objet n'est "
         "sérialisable que si tout ce qu'il contient l'est. serialVersionUID fixe la version de la classe. Le champ nbCommandesRecues est "
         "transient : il n'est pas enregistré. À droite l'écriture avec ObjectOutputStream et writeObject, puis la lecture avec "
         "readObject qui renvoie un Object : il faut le caster, et ClassNotFoundException est une exception vérifiée à traiter. "
         "En bas le fichier produit : il est binaire, commence par ac ed 00 05, la signature Java, et on y lit les noms de classes.")

# ====================================================================== 9. Q3 UML + résultats
s = prs.slides.add_slide(BLANK)
header(s, "Q3 — Résultat et UML mis à jour", 9)
picture(s, "TP5_Q3_serialisation.png", 0.25, 1.15, 6.9, 5.9, "Capture 13 — diagramme TP4 + interface Serializable (Restaurant, Plat), transient, static", folder=UML)
picture(s, "q3_console.png", 7.3, 1.2, 5.75, 5.85, "Capture 14 — avant, sérialisation, désérialisation : copie identique, transient remis à 0")
notes(s, "Q3, résultat (1 min). Notre diagramme TP4 avec l'interface Serializable ajoutée en haut à droite : Restaurant et Plat la réalisent. "
         "Dans Restaurant on voit serialVersionUID, le champ transient et les deux méthodes statiques de sauvegarde et de chargement. "
         "La note rappelle que Client, Livreur, Commande, Panier et Livraison ne sont pas sérialisables. "
         "La console : avant, deux restaurants avec leurs plats et leurs notes. Après relecture, les objets sont identiques en contenu, "
         "mais ce sont des copies, le == donne false. La note du plat est conservée et le champ transient est revenu à zéro. "
         "Deux cas d'erreur : sérialiser un Client lève NotSerializableException, et lire un fichier qui n'est pas un .ser lève "
         "StreamCorruptedException.")

# ====================================================================== 10. Bilan + sources
s = prs.slides.add_slide(BLANK)
header(s, "Bilan, sources et IA", 10)
text(s, 0.4, 1.2, 6.3, 5.8, [
    ("Bilan", 22, True, NAVY),
    ("Fait : 9 exceptions métier + hiérarchie, fichier texte (écrire, lire, afficher), sérialisation d'une liste de restaurants.", 15.5, False, DARK, True),
    ("Vérifié : TestsTP5 (tests du TP4 adaptés + nouveaux cas) : tous passent ; démonstration Main q1 / q2 / q3.", 15.5, False, DARK, True),
    ("Appris : checked / unchecked, throws et redéfinition, try-with-resources, chaînage de cause, transient, copie à la désérialisation.", 15.5, False, DARK, True),
    ("Difficultés : Panier et le throws hérité ; ne pas laisser un menu à moitié chargé.", 15.5, False, DARK, True),
    ("Limites : tableaux de 20 conservés (TP4) ; pas de « ; » dans une description ; pas de gestion de versions du .ser.", 15.5, False, DARK, True),
    ("Améliorations : listes dynamiques (ArrayList) à la place des tableaux ; format CSV/JSON standard.", 15.5, False, DARK, True)])
rect(s, 6.95, 1.2, 6.0, 5.8, LIGHT)
text(s, 7.1, 1.25, 5.75, 5.7, [
    ("Sources", 22, True, NAVY),
    ("J.-P. Segado, TP5 « Exceptions – Fichier – Sérialisation » et cours 5-6 (ECE).", 14.5, False, DARK, True),
    ("Documentation Java SE (Oracle) : Exceptions, Serializable, ObjectOutputStream.", 14.5, False, DARK, True),
    ("Notre TP3 et TP4 (Question 4 complet), diagrammes draw.io de l'équipe.", 14.5, False, DARK, True),
    ("IA utilisée", 22, True, NAVY),
    ("Claude (Anthropic) : aide à l'écriture du code, des diagrammes et des slides. "
     "Code compilé et testé ; chaque choix est expliqué dans les notes de l'oral.", 14.5, False, DARK, True),
    ("Exemples de prompts :", 15.5, True, ORANGE),
    ("« fait moi le tp 5, dis-moi ce que tu as fait et prends les captures d'écran »", 14, False, GREY, True),
    ("« à partir de TP4 question complet (fichier zip), fais les changements s'il y en a »", 14, False, GREY, True)])
notes(s, "Bilan (30 s). Nous avons traité les trois questions et vérifié avec des tests automatiques que le TP4 fonctionne toujours. "
         "Ce que nous retenons : la différence entre exceptions vérifiées et non vérifiées, la règle du throws pour les méthodes redéfinies, "
         "et le fait qu'un objet désérialisé est une copie. Limites : on garde les tableaux de 20 du TP4. "
         "Sources : le sujet et le cours de M. Segado, la documentation Java d'Oracle. Nous avons utilisé l'IA Claude pour aider à "
         "écrire le code, avec deux exemples de prompts affichés ici ; le code a été compilé et testé, et nous devons pouvoir expliquer chaque choix.")

prs.save(OUTFILE)
print("OK", OUTFILE)
